import csv
import json

import pytest
from openpyxl import load_workbook

from shopify_pipeline.exceptions import DataError
from shopify_pipeline.state import Ledger
from shopify_pipeline.transform.alertas import EVENT_COLUMNS
from shopify_pipeline.transform.run import latest_raw_file, read_jsonl, run_transform
from shopify_pipeline.transform.siesa_lines import COLUMNS
from tests.conftest import FIXTURES


def _read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


@pytest.fixture
def raw_file(tmp_path):
    order = json.loads((FIXTURES / "order_1006.json").read_text(encoding="utf-8"))
    pagado = order | {"financial_status": "paid"}
    pendiente = order | {"id": 2, "name": "#1007", "order_number": 1007}  # contraentrega sin cobrar
    path = tmp_path / "raw" / "20261014T150000Z.jsonl"
    path.parent.mkdir()
    lines = [json.dumps(o, ensure_ascii=False) for o in (pagado, pendiente)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def test_run_transform_escribe_csv_y_xlsx(raw_file, tmp_path):
    run = run_transform(raw_file, tmp_path / "out", log_dir=tmp_path / "logs")

    assert run.csv_file.name == "pedidos_siesa_20261014T150000Z.csv"
    rows = _read_csv(run.csv_file)
    assert list(rows[0]) == list(COLUMNS)
    assert [r["referencia"] for r in rows] == ["KITK0108", "ENVIO"]
    assert rows[1]["precio_unitario"] == "12294.34"

    ws = load_workbook(run.xlsx_file).active
    header = [c.value for c in ws[1]]
    precio = ws.cell(row=2, column=header.index("precio_unitario") + 1).value
    assert precio == 89999.0  # Número real en Excel, no texto
    assert ws.max_row == 3


def test_run_transform_escribe_log_de_alertas(raw_file, tmp_path, caplog):
    run = run_transform(raw_file, tmp_path / "out", log_dir=tmp_path / "logs")

    assert run.alertas_file == tmp_path / "logs" / "alertas_20261014T150000Z.csv"
    eventos = _read_csv(run.alertas_file)
    assert list(eventos[0]) == list(EVENT_COLUMNS)
    # El JSONL de prueba no trae bodegas: el #1006 sale con alerta y el #1007 queda excluido.
    assert [
        (e["tipo"], e["codigo"], e["pedido_shopify"], e["numero_linea"], e["detalle"]) for e in eventos
    ] == [
        ("ALERTA", "BODEGA_NO_CONSULTADA", "#1006", "1", ""),
        ("EXCLUIDO", "PAGO_NO_CONFIRMADO", "#1007", "", "estado_pago=pending"),
    ]
    assert "PAGO_NO_CONFIRMADO=1" in caplog.text
    assert "BODEGA_NO_CONSULTADA=1" in caplog.text


def test_log_de_alertas_vacio_igual_se_crea(tmp_path):
    raw = tmp_path / "vacio.jsonl"
    raw.write_text("", encoding="utf-8")
    run = run_transform(raw, tmp_path / "out", log_dir=tmp_path / "logs")
    assert run.alertas_file.read_text(encoding="utf-8-sig").startswith("tipo,codigo")


def test_latest_raw_file(tmp_path):
    (tmp_path / "20261013T100000Z.jsonl").write_text("", encoding="utf-8")
    (tmp_path / "20261014T100000Z.jsonl").write_text("", encoding="utf-8")
    (tmp_path / ".20261015.jsonl.tmp").write_text("", encoding="utf-8")
    assert latest_raw_file(tmp_path).name == "20261014T100000Z.jsonl"


def test_latest_raw_file_vacio(tmp_path):
    with pytest.raises(ValueError, match="extract"):
        latest_raw_file(tmp_path)


def test_read_jsonl_reporta_linea_invalida(tmp_path):
    path = tmp_path / "x.jsonl"
    path.write_text('{"id": 1}\n{malo\n', encoding="utf-8")
    with pytest.raises(DataError, match="línea 2"):
        list(read_jsonl(path))


def _manifest(raw_file, **campos):
    data = {"schema_version": 2, "reconciled": True, "fields": None} | campos
    raw_file.with_name(f"{raw_file.stem}.manifest.json").write_text(json.dumps(data), encoding="utf-8")


def test_manifiesto_con_fields_incompletos_detiene_la_transformacion(raw_file, tmp_path):
    _manifest(raw_file, fields=["id", "name", "financial_status", "line_items"])
    with pytest.raises(DataError, match="--fields"):
        run_transform(raw_file, tmp_path / "out", log_dir=tmp_path / "logs")
    assert not (tmp_path / "out").exists()  # No se generó un reporte incompleto


def test_manifiesto_sin_cuadrar_se_reporta(raw_file, tmp_path, caplog):
    _manifest(raw_file, reconciled=False, expected_count=3, written_count=2)
    run = run_transform(raw_file, tmp_path / "out", log_dir=tmp_path / "logs")
    assert not run.reconciled
    assert "no cuadró" in caplog.text


def test_manifiesto_de_version_futura(raw_file, tmp_path):
    _manifest(raw_file, schema_version=99)
    with pytest.raises(DataError, match="schema_version"):
        run_transform(raw_file, tmp_path / "out", log_dir=tmp_path / "logs")


def test_ledger_se_actualiza_despues_de_escribir(raw_file, tmp_path):
    ledger = Ledger.open(tmp_path / "state")
    run_transform(raw_file, tmp_path / "out", log_dir=tmp_path / "logs", ledger=ledger)

    reabierto = Ledger.open(tmp_path / "state")
    assert len(reabierto) == 1  # Solo el #1006 (pagado); el #1007 pendiente no se registra
    assert reabierto.get(json.loads(raw_file.read_text(encoding="utf-8").splitlines()[0])["id"]).run_id == (
        raw_file.stem
    )
