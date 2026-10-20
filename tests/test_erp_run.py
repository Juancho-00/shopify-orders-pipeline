"""Fase 3: generación de las plantillas de clientes y pedidos desde un reporte sintético."""

from __future__ import annotations

import csv
from datetime import date

import pytest
from openpyxl import load_workbook

from shopify_pipeline import cli
from shopify_pipeline.erp.run import generar_plantillas, ultimo_reporte
from shopify_pipeline.erp.terceros import TERCEROS_COLUMNS
from shopify_pipeline.exceptions import ConfigError

COLUMNAS = ["documento_referencia", "pedido_shopify", "fecha_compra", "tipo_documento", "numero_documento",
            "primer_nombre", "segundo_nombre", "primer_apellido", "segundo_apellido", "direccion",
            "direccion_complemento", "codigo_ciudad_siesa", "telefono", "email", "referencia",
            "tipo_linea", "bodega_siesa", "cantidad", "precio_neto", "notas"]  # fmt: skip


def _fila(ref, doc, fecha, **cambios):
    valores = [ref, f"#{ref}", fecha, "C", doc, "Ana", "", "Pérez", "Gómez", "Calle 1 # 2-3", "",
               "16905001", "3001112233", "ana@example.com", "SKU1", "PRODUCTO", "LCD01", "1", "89999",
               ""]  # fmt: skip
    base = dict(zip(COLUMNAS, valores, strict=True))
    return base | cambios


def _reporte(path, filas):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNAS)
        w.writeheader()
        w.writerows(filas)
    return path


@pytest.fixture
def reporte(tmp_path):
    return _reporte(
        tmp_path / "processed" / "pedidos_siesa_20261019T113000Z.csv",
        [
            _fila("1006", "1017000001", "2026-10-18 10:00:00"),
            _fila(
                "1006",
                "1017000001",
                "2026-10-18 10:00:00",
                referencia="SERV01",
                tipo_linea="ENVIO",
                precio_neto="12294",
            ),  # 2ª línea: envío
            _fila("1007", "1017000001", "2026-10-18 15:00:00", direccion="Carrera 9 # 8-7"),  # mismo cliente
            _fila("1008", "70123456", "2026-10-18 16:00:00", codigo_ciudad_siesa=""),  # sin ciudad
        ],
    )


def test_genera_un_tercero_por_cliente(reporte, tmp_path):
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 19))

    assert run.terceros_file.name == "terceros_20261019T113000Z.xlsx"
    assert run.terceros == 1  # 1006 y 1007 son del mismo cliente; 1008 no se carga
    assert run.pedidos_cargables == {"1006", "1007"}
    ws = load_workbook(run.terceros_file).active
    assert ws.title == "Terceros"
    assert [c.value for c in ws[1]] == list(TERCEROS_COLUMNS)
    fila = dict(zip(TERCEROS_COLUMNS, (c.value for c in ws[2]), strict=True))
    assert fila["F200_ID"] == 1017000001  # número, como en la plantilla
    assert fila["F200_FECHA_NACIMIENTO"] == "20261019"  # texto, como en la plantilla de ejemplo
    assert fila["F015_DIRECCION1"] == "Carrera 9 # 8-7"  # datos del pedido más reciente
    assert fila["F015_ID_DEPTO"] == "05" and ws.cell(row=2, column=14).number_format == "@"


def test_libro_tiene_las_4_hojas_de_la_plantilla(tmp_path):
    reporte = _reporte(
        tmp_path / "pedidos_siesa_20261019T113000Z.csv",
        [
            _fila("1006", "1017000001", "2026-10-18 10:00:00"),
            _fila("1007", "70123456", "2026-10-18 11:00:00"),
        ],
    )
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 19))
    wb = load_workbook(run.terceros_file)
    assert wb.sheetnames == ["Terceros", "Clientes", "Imptos y Reten", "Ent. Dinamica Cliente"]

    clientes = [[c.value for c in fila] for fila in wb["Clientes"].iter_rows(min_row=2)]
    assert [f[0] for f in clientes] == [1017000001, 70123456]
    assert clientes[0][1:8] == ["001", "PÉREZ GÓMEZ ANA", "W001", "03D", "C001", "1", "500"]
    assert clientes[0][17:19] == ["20261019", "W001"]

    impuestos = [[c.value for c in fila] for fila in wb["Imptos y Reten"].iter_rows(min_row=2)]
    assert impuestos == [  # por bloques: 46 de todos, 47/1 de todos, 47/41 de todos
        ["46", 1017000001, "001", "1", "1"],
        ["46", 70123456, "001", "1", "1"],
        ["47", 1017000001, "001", "1", "0"],
        ["47", 70123456, "001", "1", "0"],
        ["47", 1017000001, "001", "41", "1"],
        ["47", 70123456, "001", "41", "1"],
    ]
    entidad = [[c.value for c in fila] for fila in wb["Ent. Dinamica Cliente"].iter_rows(min_row=2)]
    assert entidad == [[1017000001, "001", "ana@example.com"], [70123456, "001", "ana@example.com"]]


def test_pedido_con_cliente_incompleto_queda_en_el_log(reporte, tmp_path):
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 19))
    with run.eventos_file.open(encoding="utf-8-sig") as fh:
        eventos = list(csv.DictReader(fh))
    assert [(e["codigo"], e["documento_referencia"], e["detalle"]) for e in eventos] == [
        ("TERCERO_INCOMPLETO", "1008", "falta: codigo_ciudad_siesa")
    ]
    assert "70123456" not in run.eventos_file.read_text(encoding="utf-8-sig")  # sin datos personales


def test_ultimo_reporte_ignora_revision(tmp_path):
    _reporte(tmp_path / "pedidos_siesa_20261018T113000Z.csv", [])
    _reporte(tmp_path / "pedidos_siesa_20261019T113000Z_revision.csv", [])
    assert ultimo_reporte(tmp_path).name == "pedidos_siesa_20261018T113000Z.csv"
    with pytest.raises(ConfigError):
        ultimo_reporte(tmp_path / "vacio")


def test_cli_plantillas(reporte, tmp_path):
    base = ["--log-dir", str(tmp_path / "logs"), "plantillas", "--processed-dir", str(reporte.parent),
            "--output-dir", str(tmp_path / "cargas")]  # fmt: skip
    assert cli.main(base) == cli.EXIT_OK
    assert cli.main([*base, "--strict"]) == cli.EXIT_DATA  # el 1008 no se pudo cargar
    assert (tmp_path / "cargas" / "terceros_20261019T113000Z.xlsx").exists()


def test_plantilla_de_pedidos(tmp_path):
    reporte = _reporte(
        tmp_path / "pedidos_siesa_20261020T113000Z.csv",
        [
            _fila("2001", "70123456", "2026-10-19 09:00:00", bodega_siesa="LEC04"),  # Melonn → CPM
            _fila("2002", "1017000001", "2026-10-19 10:00:00", bodega_siesa="LCM07", notas="Portería"),  # CPT
            _fila("2003", "1017000002", "2026-10-19 11:00:00", cantidad="2", precio_neto="84999"),  # CPW
            _fila(
                "2003",
                "1017000002",
                "2026-10-19 11:00:00",
                referencia="SERV01",
                tipo_linea="ENVIO",
                precio_neto="12294",
            ),
            _fila("2004", "1017000003", "2026-10-19 12:00:00"),  # CPW, segundo del bloque
        ],
    )
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 20))
    assert run.pedidos_file.name == "pedidos_20261020T113000Z.xlsx" and run.pedidos == 4
    wb = load_workbook(run.pedidos_file)
    assert wb.sheetnames == ["Pedidos", "Movto Pedidos comercial"]

    encabezados = [[c.value for c in f] for f in wb["Pedidos"].iter_rows(min_row=2)]
    # Por bloques CPW → CPT → CPM; el consecutivo se reinicia en cada bloque.
    assert [(f[1], f[2], f[11]) for f in encabezados] == [
        ("CPW", 1, "2003"),
        ("CPW", 2, "2004"),
        ("CPT", 1, "2002"),
        ("CPM", 1, "2001"),
    ]
    cpt = encabezados[2]
    assert cpt == ["001", "CPT", 1, "20261020", "4", 1017000001, "001", 1017000001, "001", "001",
                   "20261020", "2002", "03D", "Portería", "000", "900733097"]  # fmt: skip

    movtos = [[c.value for c in f] for f in wb["Movto Pedidos comercial"].iter_rows(min_row=2)]
    assert [(f[1], f[2], f[3], f[4], f[7], f[15], f[16]) for f in movtos[:2]] == [
        ("CPW", "1", 1, "SKU1", "LCD01", "2", "84999"),
        ("CPW", "1", 2, "SERV01", "LCD01", "1", "12294"),  # el envío es una línea más
    ]
    assert movtos[0][8:15] == ["06", "0", "001", None, "20261020", "03", "UND"]
    assert movtos[0][17:] == ["0", "2"]
    assert len(movtos) == 5


@pytest.mark.parametrize(
    ("cambios", "codigo"),
    [
        ({"referencia": ""}, "PEDIDO_INCOMPLETO"),
        ({"bodega_siesa": ""}, "PEDIDO_INCOMPLETO"),
        ({"bodega_siesa": "LCM54"}, "BODEGA_NO_SE_CARGA"),  # vending: no va a Siesa
        ({"bodega_siesa": "LXX99"}, "TIPO_DOCTO_SIN_EQUIVALENCIA"),  # bodega sin tipo configurado
        ({"email": ""}, "TERCERO_INCOMPLETO"),
    ],
)
def test_pedido_no_cargable_no_va_en_ninguna_plantilla(tmp_path, cambios, codigo):
    reporte = _reporte(tmp_path / "pedidos_siesa_20261020T113000Z.csv",
                       [_fila("3001", "70123456", "2026-10-19 09:00:00", **cambios)])  # fmt: skip
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 20))
    assert (run.terceros, run.pedidos) == (0, 0)  # ni el cliente ni el pedido se cargan
    assert [e.codigo for e in run.eventos] == [codigo]


def test_pedido_con_bodegas_de_tipos_distintos(tmp_path):
    reporte = _reporte(
        tmp_path / "pedidos_siesa_20261020T113000Z.csv",
        [
            _fila("3002", "70123456", "2026-10-19 09:00:00", bodega_siesa="LCD01"),
            _fila("3002", "70123456", "2026-10-19 09:00:00", referencia="SKU2", bodega_siesa="LEC04"),
        ],
    )
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 20))
    (evento,) = run.eventos
    assert (evento.codigo, evento.detalle) == ("PEDIDO_BODEGAS_MIXTAS", "LCD01=CPW / LEC04=CPM")


def test_vending_no_cuenta_como_error(tmp_path):
    reporte = _reporte(tmp_path / "pedidos_siesa_20261020T113000Z.csv",
                       [_fila("3003", "70123456", "2026-10-19 09:00:00", bodega_siesa="LCM55")])  # fmt: skip
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 20))
    assert (run.pedidos, run.vending, run.no_cargados) == (0, 1, 0)
    args = ["--log-dir", str(tmp_path / "logs"), "plantillas", "--input", str(reporte),
            "--output-dir", str(tmp_path / "cargas"), "--strict"]  # fmt: skip
    assert cli.main(args) == cli.EXIT_OK  # --strict no falla por un pedido de vending


def test_notas_se_cortan_en_500_caracteres(tmp_path):
    reporte = _reporte(tmp_path / "pedidos_siesa_20261020T113000Z.csv",
                       [_fila("3004", "70123456", "2026-10-19 09:00:00", notas="x" * 620)])  # fmt: skip
    run = generar_plantillas(reporte, tmp_path / "cargas", tmp_path / "logs", date(2026, 10, 20))
    notas = load_workbook(run.pedidos_file)["Pedidos"].cell(row=2, column=14).value
    assert len(notas) == 500
    assert [(e.codigo, e.detalle) for e in run.eventos] == [
        ("NOTAS_RECORTADAS", "620 caracteres; Siesa admite 500")
    ]
