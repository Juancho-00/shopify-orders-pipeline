"""
Escritura y lectura tabular de memoria constante: Excel en modo streaming y CSV del
exporte de Siesa leído fila por fila.
"""

from decimal import Decimal

from openpyxl import load_workbook

from shopify_pipeline.notificaciones import _filas_hoja
from shopify_pipeline.reconciliation.siesa_export import _codificacion, leer_exporte
from shopify_pipeline.sinks.tabular import FORMATO_MONEDA, write_xlsx

COLUMNAS = ("pedido", "cantidad", "precio_unitario", "descripcion")
NUMERICAS = frozenset({"cantidad", "precio_unitario"})


def _filas(n):
    return [
        {"pedido": f"#{i}", "cantidad": 2, "precio_unitario": Decimal("89999.50"), "descripcion": "Kit ñandú"}
        for i in range(n)
    ]


def test_xlsx_streaming_conserva_contenido_y_formato(tmp_path):
    path = tmp_path / "r.xlsx"
    assert write_xlsx(_filas(3), path, COLUMNAS, NUMERICAS) == 3

    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == list(COLUMNAS)
    assert ws["A1"].font.bold
    assert [c.value for c in ws[2]] == ["#0", 2, 89999.5, "Kit ñandú"]  # montos como número real
    assert ws["C2"].number_format == FORMATO_MONEDA
    assert ws["B2"].number_format == "General"  # cantidad sin formato de moneda
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:D4"


def test_xlsx_streaming_sin_filas(tmp_path):
    path = tmp_path / "vacio.xlsx"
    write_xlsx([], path, COLUMNAS, NUMERICAS)
    assert [c.value for c in load_workbook(path).active[1]] == list(COLUMNAS)


def test_conteo_de_filas_no_depende_de_la_dimension_declarada(tmp_path):
    """El modo streaming no escribe <dimension>; el aviso por correo igual cuenta bien."""
    path = tmp_path / "r.xlsx"
    write_xlsx(_filas(5), path, COLUMNAS, NUMERICAS, sheet_title="Pedidos")
    assert _filas_hoja(path, "Pedidos") == 5


# --------------------------------------------------------------- CSV de Siesa
def test_codificacion_utf8_con_caracter_partido_entre_bloques(tmp_path):
    """La 'é' (2 bytes) queda partida entre el bloque de 1 MB y el siguiente: sigue siendo UTF-8."""
    path = tmp_path / "grande.csv"
    path.write_bytes(b"a" * ((1 << 20) - 1) + "é\n".encode())
    assert _codificacion(path) == "utf-8-sig"


def test_codificacion_cp1252(tmp_path):
    path = tmp_path / "excel.csv"
    path.write_bytes("Número;_num_docto_referencia\nañ;1006\n".encode("cp1252"))
    assert _codificacion(path) == "cp1252"


def test_exporte_csv_cp1252_se_lee_en_streaming(tmp_path):
    path = tmp_path / "exporte.csv"
    lineas = ["Reporte de pedidos", "Nro documento;_num_docto_referencia;Señor"]
    lineas += [f"CPM-{i:08d};{1000 + i};x" for i in range(2000)]
    path.write_bytes(("\r\n".join(lineas) + "\r\n").encode("cp1252"))

    exporte = leer_exporte(path)

    assert len(exporte.referencias) == 2000
    assert "1006" in exporte.referencias
    assert exporte.documentos["1006"] == "CPM-00000006"
