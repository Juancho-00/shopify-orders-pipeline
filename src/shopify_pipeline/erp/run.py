"""
Fase 3: genera las plantillas de carga a Siesa a partir del reporte de la fase 1.

QUÉ HACE (plantilla de clientes con sus 4 hojas y plantilla de pedidos con sus 2 hojas)
    1. Lee un reporte oficial ``pedidos_siesa_<run_id>.csv`` (los ``*_revision`` no se
       cargan a Siesa).
    2. Arma un tercero por pedido con ``terceros.fila_tercero`` y, a partir de él, las hojas
       ``Clientes``, ``Imptos y Reten`` y ``Ent. Dinamica Cliente`` (``clientes.py``). El
       cliente se sube con cada pedido aunque ya exista (R-31); si el mismo cliente compró
       varias veces en la corrida, sale una sola vez con los datos de su pedido más reciente.
    3. Arma el encabezado y el detalle de cada pedido (``pedidos.py``): tipo de documento
       según la bodega, consecutivo por bloques CPW → CPT → CPM.
    4. Si a un pedido le falta un dato obligatorio, del cliente o del pedido, no va en
       NINGUNA de las dos plantillas (R-34) y queda en el log con el motivo:
       ``TERCERO_INCOMPLETO``, ``PEDIDO_INCOMPLETO``, ``TIPO_DOCTO_SIN_EQUIVALENCIA`` o
       ``PEDIDO_BODEGAS_MIXTAS``. Así un cliente nunca queda cargado sin su pedido ni al revés.
       Los pedidos de vending tampoco van (``BODEGA_NO_SE_CARGA``), pero no son un error.

SALIDA (orden de carga: primero clientes, luego pedidos; R-30)
    data/siesa/cargas/terceros_<run_id>.xlsx        # 1. plantilla de clientes
    data/siesa/cargas/pedidos_<run_id>.xlsx         # 2. plantilla de pedidos
    data/logs/plantillas_<run_id>.csv               # pedidos que no se cargan y alertas
"""

from __future__ import annotations

import csv
import logging
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from shopify_pipeline.erp.clientes import (
    CLIENTES_COLUMNS,
    ENTIDAD_COLUMNS,
    IMPUESTOS_COLUMNS,
    fila_cliente,
    fila_entidad,
    filas_impuestos,
)
from shopify_pipeline.erp.pedidos import (
    LARGO_NOTAS,
    MOVTO_COLUMNS,
    PEDIDOS_COLUMNS,
    fila_pedido,
    filas_movimiento,
    numerar,
    validar_pedido,
)
from shopify_pipeline.erp.terceros import COLUMNAS_NUMERICAS, TERCEROS_COLUMNS, fila_tercero
from shopify_pipeline.exceptions import ConfigError
from shopify_pipeline.sinks.atomic import atomic_target
from shopify_pipeline.sinks.tabular import write_csv
from shopify_pipeline.transform.alertas import EVENT_COLUMNS, Evento

logger = logging.getLogger(__name__)

DEFAULT_CARGAS_DIR = Path("data/siesa/cargas")


@dataclass(frozen=True)
class PlantillasRun:
    terceros_file: Path
    pedidos_file: Path
    eventos_file: Path
    terceros: int
    pedidos: int
    pedidos_cargables: frozenset[str]  # documento_referencia de los pedidos que sí se suben
    eventos: tuple[Evento, ...]

    @property
    def no_cargados(self) -> int:
        """Pedidos que debían ir a Siesa y no van por datos incompletos (sin contar vending)."""
        return sum(1 for e in self.eventos if e.tipo == "EXCLUIDO" and e.codigo != "BODEGA_NO_SE_CARGA")

    @property
    def vending(self) -> int:
        """Pedidos de bodegas de vending: no se cargan a Siesa por definición."""
        return sum(1 for e in self.eventos if e.codigo == "BODEGA_NO_SE_CARGA")


def ultimo_reporte(processed_dir: Path) -> Path:
    """El reporte oficial más reciente (el ``run_id`` es ordenable)."""
    reportes = sorted(
        p for p in processed_dir.glob("pedidos_siesa_*.csv") if not p.stem.endswith("_revision")
    )
    if not reportes:
        raise ConfigError(f"No hay reportes de la fase 1 (pedidos_siesa_*.csv) en {processed_dir}.")
    return reportes[-1]


def _pedidos(reporte: Path) -> dict[str, list[dict[str, str]]]:
    """Líneas del reporte agrupadas por pedido (``documento_referencia``), en su orden."""
    pedidos: dict[str, list[dict[str, str]]] = {}
    with reporte.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            ref = (row.get("documento_referencia") or "").strip()
            if ref:
                pedidos.setdefault(ref, []).append(row)
    return pedidos


# Hoja → (columnas, columnas que van como número). Igual que la plantilla de ejemplo.
HOJAS: dict[str, tuple[tuple[str, ...], frozenset[str]]] = {
    # Plantilla de clientes
    "Terceros": (TERCEROS_COLUMNS, COLUMNAS_NUMERICAS),
    "Clientes": (CLIENTES_COLUMNS, frozenset({"F201_ID_TERCERO"})),
    "Imptos y Reten": (IMPUESTOS_COLUMNS, frozenset({"F_ID_TERCERO"})),
    "Ent. Dinamica Cliente": (ENTIDAD_COLUMNS, frozenset({"f201_id_tercero"})),
    # Plantilla de pedidos
    "Pedidos": (PEDIDOS_COLUMNS, frozenset({"_consec_docto", "_id_tercero_fact", "_id_tercero_rem"})),
    "Movto Pedidos comercial": (MOVTO_COLUMNS, frozenset({"_nro_registro"})),
}


def _escribir_xlsx(hojas: dict[str, list[dict[str, str]]], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for titulo, filas in hojas.items():
        columnas, numericas = HOJAS[titulo]
        ws = wb.create_sheet(titulo)
        ws.append(list(columnas))
        for celda in ws[1]:
            celda.font = Font(bold=True)
        for fila in filas:
            ws.append([int(fila[c]) if c in numericas and fila[c].isdigit() else fila[c] for c in columnas])
        # Texto (@) para conservar ceros a la izquierda: 001, 05, 03D, fechas AAAAMMDD...
        for fila_celdas in ws.iter_rows(min_row=2):
            for celda in fila_celdas:
                if isinstance(celda.value, str):
                    celda.number_format = "@"
        for i, col in enumerate(columnas, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(col) + 2)
        ws.freeze_panes = "A2"
    with atomic_target(path) as tmp:
        wb.save(tmp)


def generar_plantillas(
    reporte: Path,
    output_dir: Path = DEFAULT_CARGAS_DIR,
    log_dir: Path = Path("data/logs"),
    fecha_carga: date | None = None,
) -> PlantillasRun:
    run_id = reporte.stem.removeprefix("pedidos_siesa_")
    fecha = (fecha_carga or date.today()).strftime("%Y%m%d")

    eventos: list[Evento] = []
    cargables: list[tuple[str, list[dict[str, str]], dict[str, str]]] = []  # (tipo, líneas, tercero)
    for ref, lineas in _pedidos(reporte).items():
        nombre = lineas[0].get("pedido_shopify", "")
        tercero = fila_tercero(lineas[0], fecha)
        validacion = validar_pedido(lineas)
        if not tercero.cargable:
            # Solo el nombre del campo, nunca el dato del cliente.
            detalle = "falta: " + ", ".join(tercero.faltantes)
            eventos.append(Evento("EXCLUIDO", "TERCERO_INCOMPLETO", nombre, ref, detalle=detalle))
        if not validacion.cargable:
            eventos.append(Evento("EXCLUIDO", validacion.motivo, nombre, ref, detalle=validacion.detalle))
        if not (tercero.cargable and validacion.cargable):
            continue
        for codigo in tercero.alertas:
            eventos.append(Evento("ALERTA", codigo, nombre, ref))
        largo_notas = len((lineas[0].get("notas") or "").strip())
        if largo_notas > LARGO_NOTAS:
            detalle = f"{largo_notas} caracteres; Siesa admite {LARGO_NOTAS}"
            eventos.append(Evento("ALERTA", "NOTAS_RECORTADAS", nombre, ref, detalle=detalle))
        cargables.append((validacion.tipo_docto, lineas, tercero.fila))

    # Clientes: uno por documento, con los datos de su pedido más reciente (R-31).
    por_documento: dict[str, tuple[str, dict[str, str], str]] = {}
    for _, lineas, tercero in cargables:
        fecha_compra = lineas[0].get("fecha_compra", "")
        anterior = por_documento.get(tercero["F200_ID"])
        if anterior is None or fecha_compra >= anterior[0]:
            por_documento[tercero["F200_ID"]] = (fecha_compra, tercero, lineas[0]["documento_referencia"])
    terceros = [fila for _, fila, _ in por_documento.values()]
    clientes = []
    for _, tercero, ref in por_documento.values():
        cliente, recortada = fila_cliente(tercero, fecha)
        clientes.append(cliente)
        if recortada:
            eventos.append(Evento("ALERTA", "DIRECCION_TRUNCADA", f"#{ref}", ref, detalle="hoja Clientes"))

    # Pedidos: encabezado + detalle, numerados por bloques de tipo de documento.
    documento_de = {lineas[0]["documento_referencia"]: t["F200_ID"] for _, lineas, t in cargables}
    encabezados, movimientos = [], []
    for tipo, consec, lineas in numerar([(tipo, lineas) for tipo, lineas, _ in cargables]):
        documento = documento_de[lineas[0]["documento_referencia"]]
        encabezados.append(fila_pedido(lineas, tipo, consec, fecha, documento))
        movimientos.extend(filas_movimiento(lineas, tipo, consec, fecha))

    terceros_file = output_dir / f"terceros_{run_id}.xlsx"
    pedidos_file = output_dir / f"pedidos_{run_id}.xlsx"
    eventos_file = log_dir / f"plantillas_{run_id}.csv"
    _escribir_xlsx(
        {
            "Terceros": terceros,
            "Clientes": clientes,
            "Imptos y Reten": filas_impuestos([f["F200_ID"] for f in terceros]),
            "Ent. Dinamica Cliente": [fila_entidad(f) for f in terceros],
        },
        terceros_file,
    )
    _escribir_xlsx({"Pedidos": encabezados, "Movto Pedidos comercial": movimientos}, pedidos_file)
    write_csv((e.as_row() for e in eventos), eventos_file, EVENT_COLUMNS)

    run = PlantillasRun(
        terceros_file,
        pedidos_file,
        eventos_file,
        len(terceros),
        len(encabezados),
        frozenset(documento_de),
        tuple(eventos),
    )
    logger.info("Plantilla de clientes: %d clientes (%s).", len(terceros), terceros_file)
    logger.info(
        "Plantilla de pedidos: %d pedidos, %d líneas (%s).", len(encabezados), len(movimientos), pedidos_file
    )
    if run.vending:
        logger.info("%d pedidos de vending no se cargan a Siesa (por definición).", run.vending)
    if run.no_cargados:
        logger.warning(
            "%d pedidos NO se cargan por datos incompletos. Detalle en %s",
            run.no_cargados,
            eventos_file,
        )
    alertas = Counter(e.codigo for e in eventos if e.tipo == "ALERTA")
    if alertas:
        logger.warning("Alertas: %s.", ", ".join(f"{c}={n}" for c, n in alertas.most_common()))
    return run
