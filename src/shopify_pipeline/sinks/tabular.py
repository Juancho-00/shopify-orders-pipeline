"""
Sinks tabulares: CSV (para sistemas) y Excel (para revisión humana).

* CSV: UTF-8 con BOM, separador ``,`` y punto decimal. Es el formato estable para
  cargas automáticas.
* Excel: los montos se guardan como números reales, así que se ven bien sin importar
  la configuración regional del equipo (en Colombia, Excel usa coma decimal y abriría
  mal un CSV con punto decimal).

MEMORIA
    El Excel se escribe en modo *streaming* (``Workbook(write_only=True)``): cada fila se
    vuelca al archivo al agregarla, en vez de mantener un objeto por celda. Con 6.000 filas
    el pico de memoria baja de ~68 MB a ~0,5 MB, con el mismo contenido y casi el mismo
    tiempo. Por eso todo el formato (anchos, panel congelado, formato numérico) se define
    ANTES o MIENTRAS se escriben las filas: en este modo no se puede volver atrás.

Ambos se escriben de forma atómica (temporal + renombrar).
"""

from __future__ import annotations

import csv
from collections.abc import Iterable, Sequence
from decimal import Decimal
from pathlib import Path
from typing import Any

from shopify_pipeline.sinks.atomic import atomic_target

# Columnas numéricas que NO llevan formato de moneda.
_SIN_FORMATO_MONEDA = frozenset({"cantidad", "tasa_iva"})
FORMATO_MONEDA = "#,##0.00"


def write_csv(rows: Iterable[dict[str, Any]], path: Path, columns: Sequence[str]) -> int:
    count = 0
    with atomic_target(path) as tmp, tmp.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
            count += 1
    return count


def write_xlsx(
    rows: Sequence[dict[str, Any]],
    path: Path,
    columns: Sequence[str],
    numeric_columns: frozenset[str] = frozenset(),
    sheet_title: str = "Pedidos",
) -> int:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_title)
    # Anchos según las primeras 200 filas (basta para ver el contenido típico).
    for idx, col in enumerate(columns, start=1):
        width = max([len(col)] + [len(str(r.get(col, ""))) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(width, 50)
    ws.freeze_panes = "A2"

    negrita = Font(bold=True)
    encabezado = []
    for col in columns:
        celda = WriteOnlyCell(ws, value=col)
        celda.font = negrita
        encabezado.append(celda)
    ws.append(encabezado)

    moneda = {c for c in columns if c in numeric_columns and c not in _SIN_FORMATO_MONEDA}
    for row in rows:
        values: list[Any] = []
        for col in columns:
            value = row.get(col, "")
            if col in numeric_columns and isinstance(value, Decimal):
                value = float(value)
            if col in moneda:
                celda = WriteOnlyCell(ws, value=value)
                celda.number_format = FORMATO_MONEDA
                value = celda
            values.append(value)
        ws.append(values)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    with atomic_target(path) as tmp:
        wb.save(tmp)
    return len(rows)
