"""
Lectura del exporte de pedidos de Siesa (Excel o CSV).

QUÉ HACE
    De un archivo exportado desde Siesa saca el conjunto de números de pedido Shopify que
    ya existen en el ERP. Siesa los guarda en el campo ``_num_docto_referencia`` (R-21).

POR QUÉ ES TOLERANTE
    Los exportes de Siesa cambian según quién los saque: a veces traen filas de título
    antes de los encabezados, el nombre de la columna puede venir con o sin ``_`` o en
    mayúsculas, y Excel convierte ``1006`` en ``1006.0``. Aquí se busca la fila de
    encabezados en las primeras filas y se normalizan los valores para que el cruce no
    falle por formato.
"""

from __future__ import annotations

import codecs
import csv
import re
import unicodedata
from collections import Counter
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from shopify_pipeline.exceptions import ConfigError, DataError

# Nombres aceptados para la columna de referencia, ya normalizados (ver ``_clave``).
COLUMNAS_REFERENCIA = ("numdoctoreferencia", "numdocumentoreferencia", "doctoreferencia")
# Columna opcional con el número del pedido en Siesa: tipo de documento + consecutivo
# (``CPM-00003210``). Si viene, se muestra en la conciliación.
COLUMNAS_DOCUMENTO_SIESA = ("nrodocumento", "nrodocto", "numerodocumento", "documento")
FILAS_BUSQUEDA_ENCABEZADO = 20
EXTENSIONES = (".xlsx", ".xlsm", ".csv")


def _clave(nombre: Any) -> str:
    """'_num_docto_referencia' / 'Num Docto Referencia' → 'numdoctoreferencia'."""
    texto = unicodedata.normalize("NFKD", str(nombre or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def normalizar_referencia(valor: Any) -> str:
    """
    Deja la referencia comparable con ``documento_referencia`` del reporte.

    ``' #1006 '`` → ``'1006'``; ``1006.0`` (número de Excel) → ``'1006'``.
    """
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip().lstrip("#").strip()
    return re.sub(r"\.0+$", "", texto)


@dataclass(frozen=True)
class ExporteSiesa:
    archivo: Path
    columna: str
    referencias: frozenset[str]
    duplicadas: dict[str, int] = field(default_factory=dict)
    # Referencia Shopify → número(s) del pedido en Siesa (``CPM-00003210``), si el exporte lo trae.
    documentos: dict[str, str] = field(default_factory=dict)


def _filas_xlsx(path: Path) -> Iterator[Sequence[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        yield from wb.worksheets[0].iter_rows(values_only=True)  # El exporte trae una sola hoja.
    finally:
        wb.close()


def _codificacion(path: Path) -> str:
    """UTF-8 (con o sin BOM) o, si no decodifica, cp1252 (Excel en español)."""
    # Decodificador incremental: valida por bloques de 1 MB sin cargar el archivo, y un
    # carácter partido entre dos bloques no cuenta como error.
    decoder = codecs.getincrementaldecoder("utf-8")()
    with path.open("rb") as fh:
        try:
            for bloque in iter(lambda: fh.read(1 << 20), b""):
                decoder.decode(bloque)
            decoder.decode(b"", final=True)
        except UnicodeDecodeError:
            return "cp1252"
    return "utf-8-sig"


def _filas_csv(path: Path) -> Iterator[Sequence[Any]]:
    """
    Lee el CSV fila por fila (memoria constante). Antes se cargaba el archivo completo y se
    partía en líneas: dos copias en memoria del exporte, que es el archivo más grande.
    """
    encoding = _codificacion(path)
    with path.open(encoding=encoding, newline="") as fh:
        muestra = fh.read(4096)
        fh.seek(0)
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
        except csv.Error:
            dialecto = csv.excel
        yield from csv.reader(fh, dialecto)


@dataclass(frozen=True)
class _Encabezado:
    indice_ref: int
    nombre_ref: str
    indice_doc: int | None


def _buscar_encabezado(filas: Iterator[Sequence[Any]], columna: str | None) -> _Encabezado:
    buscadas = (_clave(columna),) if columna else COLUMNAS_REFERENCIA
    for _ in range(FILAS_BUSQUEDA_ENCABEZADO):
        fila = next(filas, None)
        if fila is None:
            break
        claves = [_clave(c) for c in fila]
        for buscada in buscadas:
            if buscada in claves:
                indice = claves.index(buscada)
                doc = next((claves.index(c) for c in COLUMNAS_DOCUMENTO_SIESA if c in claves), None)
                return _Encabezado(indice, str(fila[indice]), doc)
    raise DataError(
        f"No se encontró la columna {columna or '_num_docto_referencia'} en las primeras "
        f"{FILAS_BUSQUEDA_ENCABEZADO} filas del exporte de Siesa."
    )


def _celda(fila: Sequence[Any], indice: int | None) -> Any:
    return fila[indice] if indice is not None and indice < len(fila) else None


def leer_exporte(path: Path, columna: str | None = None) -> ExporteSiesa:
    """Lee un exporte de Siesa y devuelve las referencias (números de pedido) encontradas."""
    if path.suffix.lower() not in EXTENSIONES:
        raise ConfigError(f"Formato no soportado: {path.name} (use .xlsx o .csv).")
    filas = iter(_filas_csv(path) if path.suffix.lower() == ".csv" else _filas_xlsx(path))
    enc = _buscar_encabezado(filas, columna)
    conteo: Counter[str] = Counter()
    documentos: dict[str, list[str]] = {}
    for fila in filas:
        ref = normalizar_referencia(_celda(fila, enc.indice_ref))
        if not ref:
            continue
        conteo[ref] += 1
        doc = str(_celda(fila, enc.indice_doc) or "").strip()
        if doc:
            documentos.setdefault(ref, []).append(doc)
    duplicadas = {ref: n for ref, n in conteo.items() if n > 1}
    return ExporteSiesa(
        path,
        enc.nombre_ref,
        frozenset(conteo),
        duplicadas,
        {ref: ", ".join(docs) for ref, docs in documentos.items()},
    )


def ultimo_exporte(carpeta: Path) -> Path:
    """El exporte más reciente de la carpeta (por fecha de modificación)."""
    if not carpeta.is_dir():
        # Primera vez: se crea la carpeta para que la persona sepa dónde dejar el exporte.
        carpeta.mkdir(parents=True, exist_ok=True)
        raise ConfigError(
            f"Guarde el exporte de pedidos de Siesa (.xlsx o .csv) en {carpeta} y vuelva a ejecutar."
        )
    archivos = [
        p for p in carpeta.iterdir() if p.suffix.lower() in EXTENSIONES and not p.name.startswith(("~$", "."))
    ]
    if not archivos:
        raise ConfigError(f"No hay exportes de Siesa (.xlsx/.csv) en {carpeta}.")
    return max(archivos, key=lambda p: p.stat().st_mtime)
